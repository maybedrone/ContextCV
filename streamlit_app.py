"""Streamlit frontend for the Job Application Assistant.

Flow:  input (validated) -> run to pause -> review (HITL) -> approve -> done
Writer returns STRUCTURED output (dict), enabling formatted .docx export.
"""
import io
import uuid
import streamlit as st
from pypdf import PdfReader
from docx import Document

from src.graph import app

st.set_page_config(page_title="Job Application Assistant", page_icon="📄", layout="wide")
st.title("📄 Job Application Assistant")
st.caption("Multi-agent CV tailoring with company research and human review.")

MIN_CHARS = 100


# ── HELPERS ───────────────────────────────────────────────────────────────────
def extract_pdf_text(uploaded_file) -> str:
    try:
        reader = PdfReader(uploaded_file)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:
        st.error(f"Could not read PDF: {e}")
        return ""


def cv_to_text(cv: dict) -> str:
    """Flatten the structured CV dict into readable plain text (for display + .txt)."""
    if not isinstance(cv, dict):
        return str(cv)
    lines = [cv.get("name", ""), cv.get("contact", ""), "", "SUMMARY", cv.get("summary", ""), ""]
    for section in cv.get("sections", []):
        lines.append(section.get("heading", "").upper())
        for bullet in section.get("bullets", []):
            lines.append(f"- {bullet}")
        lines.append("")
    return "\n".join(lines)


def make_docx(cv: dict) -> bytes:
    """Map the structured CV dict to a formatted .docx (bold headings, real bullets)."""
    doc = Document()
    doc.add_heading(cv.get("name", "Tailored CV"), level=0)
    if cv.get("contact"):
        doc.add_paragraph(cv["contact"])
    if cv.get("summary"):
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(cv["summary"])
    for section in cv.get("sections", []):
        doc.add_heading(section.get("heading", ""), level=1)
        for bullet in section.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def validate(jd: str, cv: str):
    problems = []
    if len(jd.strip()) < MIN_CHARS:
        problems.append(f"Job description looks too short (min {MIN_CHARS} characters).")
    if len(cv.strip()) < MIN_CHARS:
        problems.append(f"CV looks too short (min {MIN_CHARS} characters).")
    return problems


def new_session():
    st.session_state.stage = "input"
    st.session_state.thread_id = str(uuid.uuid4())


# ── SESSION STATE ─────────────────────────────────────────────────────────────
if "stage" not in st.session_state:
    new_session()

config = {
    "configurable": {"thread_id": st.session_state.thread_id},
    "recursion_limit": 15,
}

# ── STAGE 1: INPUT ────────────────────────────────────────────────────────────
if st.session_state.stage == "input":
    st.subheader("1. Provide the job description and your CV")

    jd_mode = st.radio("Job Description input:", ["Paste text", "Upload PDF"], key="jd_mode", horizontal=True)
    jd = ""
    if jd_mode == "Paste text":
        jd = st.text_area("Job Description", height=220, placeholder="Paste the full job description here...")
    else:
        jd_file = st.file_uploader("Upload the JD (PDF)", type=["pdf"], key="jd_pdf")
        if jd_file:
            jd = extract_pdf_text(jd_file)
            with st.expander("Preview extracted JD text"):
                st.text(jd)

    cv_mode = st.radio("CV input:", ["Paste text", "Upload PDF"], key="cv_mode", horizontal=True)
    cv = ""
    if cv_mode == "Paste text":
        cv = st.text_area("Your CV", height=220, placeholder="Paste your CV text here...")
    else:
        cv_file = st.file_uploader("Upload your CV (PDF)", type=["pdf"], key="cv_pdf")
        if cv_file:
            cv = extract_pdf_text(cv_file)
            with st.expander("Preview extracted CV text"):
                st.text(cv)

    if st.button("Run Assistant", type="primary"):
        problems = validate(jd, cv)
        if problems:
            for p in problems:
                st.error(p)
        else:
            initial_state = {
                "messages": [], "jd": jd, "cv": cv,
                "research": "", "analysis": "", "cv_tailored": "", "next": "",
            }
            try:
                with st.spinner("Researching company and analyzing your CV..."):
                    app.invoke(initial_state, config)
                st.session_state.stage = "review"
                st.rerun()
            except Exception as e:
                st.error("The assistant hit an error during research/analysis "
                         "(free-tier model tool-calling can be flaky). Please try again.")
                st.caption(f"Details: {e}")
                new_session()

# ── STAGE 2: REVIEW (HITL) ────────────────────────────────────────────────────
elif st.session_state.stage == "review":
    st.subheader("2. Review the research and analysis")

    try:
        state = app.get_state(config)
        with st.expander("Company Research", expanded=True):
            st.write(state.values.get("research", "No research produced."))
        with st.expander("CV vs JD Analysis", expanded=True):
            st.write(state.values.get("analysis", "No analysis produced."))
    except Exception as e:
        st.error("Could not load the run state. Please start over.")
        st.caption(f"Details: {e}")

    st.info("Approve to let the writer generate your tailored CV.")

    col1, col2 = st.columns(2)
    with col1:
        if st.button("Approve & Generate CV", type="primary"):
            try:
                with st.spinner("Writing your tailored CV..."):
                    app.invoke(None, config)
                st.session_state.stage = "done"
                st.rerun()
            except Exception as e:
                st.error("The writer hit an error. Try approving again, or start over.")
                st.caption(f"Details: {e}")
    with col2:
        if st.button("Start Over"):
            new_session()
            st.rerun()

# ── STAGE 3: DONE ─────────────────────────────────────────────────────────────
elif st.session_state.stage == "done":
    st.subheader("3. Your tailored CV")

    try:
        state = app.get_state(config)
        cv_dict = state.values.get("cv_tailored", {})
    except Exception as e:
        cv_dict = {}
        st.error("Could not load the tailored CV. Please start over.")
        st.caption(f"Details: {e}")

    readable = cv_to_text(cv_dict)
    st.text_area("Tailored CV", value=readable, height=500)

    col1, col2 = st.columns(2)
    with col1:
        st.download_button(
            label="Download as .txt",
            data=readable,
            file_name="tailored_cv.txt",
            mime="text/plain",
        )
    with col2:
        if cv_dict:
            st.download_button(
                label="Download as .docx",
                data=make_docx(cv_dict),
                file_name="tailored_cv.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

    if st.button("Start Over"):
        new_session()
        st.rerun()